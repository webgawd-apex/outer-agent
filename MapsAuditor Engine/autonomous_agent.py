"""
autonomous_agent.py
======================================================
Render-Optimized Master Autonomous Lead Finder & Outreach Agent.
Features:
  - Low-memory Chromium configuration (under 512MB RAM for Render Free tier)
  - Asset-blocking router (skips images/fonts/media on target sites to save memory & speed up scans)
  - 10 Niches x 10 Cities iteration (100 search combinations)
  - Contact form autofill + submit with anti-bot/CAPTCHA detection
  - Direct email fallback via Gmail SMTP or SendGrid
  - Enforces hard 200/day limit
  - Real-time Discord alerts + SQLite CRM logging
"""

import os
import sys
import time
import json
import random
import datetime
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright, Page, TimeoutError as PlaywrightTimeoutError

import notifier
import email_finder
import email_sender
import form_submitter
import followup_crm

SETTINGS_PATH = os.path.join(os.path.dirname(__file__), "settings.json")
EXCEL_PATH = os.path.join(os.path.dirname(__file__), "phoenix_lead_tracker.xlsx")
DAILY_LOG_PATH = os.path.join(os.path.dirname(__file__), "daily_stats.json")

def load_settings():
    settings = {}
    if os.path.exists(SETTINGS_PATH):
        try:
            with open(SETTINGS_PATH, "r", encoding="utf-8") as f:
                settings = json.load(f)
        except Exception as e:
            print(f"[agent] Error reading settings.json: {e}")
            
    # Allow Environment Variable overrides for Render deployment
    if "DISCORD_WEBHOOK_URL" in os.environ:
        settings.setdefault("notifications", {})["discord_webhook_url"] = os.environ["DISCORD_WEBHOOK_URL"]
    if "GMAIL_EMAIL" in os.environ:
        settings.setdefault("gmail", {})["email"] = os.environ["GMAIL_EMAIL"]
    if "GMAIL_APP_PASSWORD" in os.environ:
        settings.setdefault("gmail", {})["app_password"] = os.environ["GMAIL_APP_PASSWORD"]
    if "SENDGRID_API_KEY" in os.environ:
        settings.setdefault("sendgrid", {})["api_key"] = os.environ["SENDGRID_API_KEY"]
        
    return settings

def get_today_action_count() -> int:
    """Read today's count of actions (forms + emails) sent."""
    today_str = datetime.date.today().isoformat()
    if os.path.exists(DAILY_LOG_PATH):
        try:
            with open(DAILY_LOG_PATH, "r", encoding="utf-8") as f:
                data = json.load(f)
                if data.get("date") == today_str:
                    return data.get("count", 0)
        except Exception:
            pass
    return 0

def increment_today_action_count() -> int:
    """Increment and save today's action count."""
    today_str = datetime.date.today().isoformat()
    current = get_today_action_count() + 1
    try:
        with open(DAILY_LOG_PATH, "w", encoding="utf-8") as f:
            json.dump({"date": today_str, "count": current}, f, indent=2)
    except Exception as e:
        print(f"[agent] Failed to save daily stats: {e}")
    return current

def init_excel_tracker():
    """Ensure Excel tracker exists with proper headers."""
    if not os.path.exists(EXCEL_PATH):
        wb = Workbook()
        ws = wb.active
        ws.title = "Outreach Tracker"
        headers = [
            "Timestamp", "Business Name", "Niche", "City", "Phone", 
            "Rating", "Reviews", "Website", "Email", "Status", "Notes"
        ]
        ws.append(headers)
        wb.save(EXCEL_PATH)

def log_to_excel(business_name, niche, city, phone, rating, reviews, website, email, status, notes=""):
    """Append row to Excel tracker."""
    init_excel_tracker()
    try:
        wb = load_workbook(EXCEL_PATH)
        ws = wb.active
        ws.append([
            datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            business_name,
            niche,
            city,
            phone,
            rating,
            reviews,
            website,
            email,
            status,
            notes
        ])
        wb.save(EXCEL_PATH)
    except Exception as e:
        print(f"[agent] Excel log error: {e}")

def run_maps_autonomous(headless: bool = True):
    settings = load_settings()
    niches = settings.get("targets", {}).get("niches", ["Roofing"])
    cities = settings.get("targets", {}).get("cities", ["Phoenix"])
    max_daily_limit = settings.get("pipeline", {}).get("daily_outreach_limit", 200)
    delay_seconds = settings.get("pipeline", {}).get("email_delay_seconds", 8)

    crm = followup_crm.CRM()
    init_excel_tracker()

    print("\n" + "="*60)
    print(" 🚀 MAPSAUDITOR AUTONOMOUS AGENT (RENDER OPTIMIZED)")
    print(f" Targets: {len(niches)} Niches x {len(cities)} Cities = {len(niches)*len(cities)} Combinations")
    print(f" Daily Action Limit: {max_daily_limit} (Forms + Emails)")
    print(f" Current Today Count: {get_today_action_count()}/{max_daily_limit}")
    print("="*60 + "\n")

    # Render Memory-Optimized Chromium Flags (limits RAM consumption under 512MB)
    chromium_args = [
        "--no-sandbox",
        "--disable-setuid-sandbox",
        "--disable-dev-shm-usage",
        "--disable-gpu",
        "--disable-software-rasterizer",
        "--no-zygote",
        "--single-process",
        "--mute-audio",
        "--disable-background-networking",
        "--disable-default-apps",
        "--disable-extensions",
        "--disable-sync",
        "--js-flags=--max-old-space-size=256"
    ]

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=headless,
            args=chromium_args
        )
        context = browser.new_context(
            viewport={"width": 1280, "height": 800},
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
        )
        maps_page = context.new_page()

        for niche in niches:
            for city in cities:
                today_count = get_today_action_count()
                if today_count >= max_daily_limit:
                    print(f"\n🛑 Daily limit reached ({today_count}/{max_daily_limit}). Pausing engine.")
                    notifier.notify_daily_limit_reached(today_count, max_daily_limit)
                    browser.close()
                    return

                query = f"{niche} in {city}, AZ"
                print(f"\n🔍 Searching Google Maps for: '{query}'")

                try:
                    maps_page.goto("https://www.google.com/maps", timeout=30000, wait_until="domcontentloaded")
                    time.sleep(2)

                    # Handle Google consent popup if present
                    try:
                        consent_btn = maps_page.locator("button:has-text('Accept all'), button:has-text('I agree')").first
                        if consent_btn.is_visible(timeout=1500):
                            consent_btn.click()
                            time.sleep(1)
                    except Exception:
                        pass

                    # Type search query
                    search_box = maps_page.locator("#searchboxinput")
                    search_box.wait_for(timeout=10000)
                    search_box.fill(query)
                    maps_page.keyboard.press("Enter")
                    time.sleep(4)

                    # Wait for results panel
                    maps_page.wait_for_selector("div[role='feed'], div[aria-label*='Results for']", timeout=15000)

                    # Scroll the left results feed to load listings
                    feed = maps_page.locator("div[role='feed']").first
                    for _ in range(3):
                        if feed.is_visible():
                            feed.evaluate("el => el.scrollTop += 800")
                            time.sleep(1.5)

                    # Find listing cards
                    cards = maps_page.locator("div[role='feed'] > div > div > a[href*='/maps/place/']").all()
                    if not cards:
                        cards = maps_page.locator("a[href*='/maps/place/']").all()

                    print(f"📋 Found {len(cards)} listings for '{query}'")

                    for idx, card in enumerate(cards):
                        today_count = get_today_action_count()
                        if today_count >= max_daily_limit:
                            print(f"\n🛑 Daily limit reached ({today_count}/{max_daily_limit}). Stopping.")
                            notifier.notify_daily_limit_reached(today_count, max_daily_limit)
                            browser.close()
                            return

                        try:
                            # Scroll card into view and click
                            card.scroll_into_view_if_needed(timeout=3000)
                            card.click(timeout=5000)
                            time.sleep(2.5)

                            # Extract detail panel info
                            biz_name = ""
                            try:
                                biz_name = maps_page.locator("h1.DUwDvf, h1.fontHeadlineLarge, div.fontHeadlineLarge").first.inner_text(timeout=3000).strip()
                            except Exception:
                                biz_name = f"Business {idx+1}"

                            rating_str = ""
                            review_count_str = ""
                            try:
                                rating_el = maps_page.locator("div.F7nice span[aria-hidden='true']").first
                                if rating_el.is_visible(timeout=1000):
                                    rating_str = rating_el.inner_text().strip()
                                reviews_el = maps_page.locator("div.F7nice span span[aria-label*='reviews']").first
                                if reviews_el.is_visible(timeout=1000):
                                    review_count_str = reviews_el.inner_text().replace("(", "").replace(")", "").strip()
                            except Exception:
                                pass

                            phone = ""
                            try:
                                phone_btn = maps_page.locator("button[data-tooltip*='Copy phone number'], button[aria-label*='Phone:']").first
                                if phone_btn.is_visible(timeout=1000):
                                    phone = phone_btn.get_attribute("aria-label").replace("Phone:", "").strip()
                            except Exception:
                                pass

                            website = ""
                            try:
                                web_btn = maps_page.locator("a[data-tooltip*='Open website'], a[aria-label*='Website:']").first
                                if web_btn.is_visible(timeout=1500):
                                    website = web_btn.get_attribute("href")
                            except Exception:
                                pass

                            print(f"\n[{idx+1}/{len(cards)}] 🏢 {biz_name} | ⭐ {rating_str} ({review_count_str} rev) | 🌐 {website or 'NO WEBSITE'}")

                            # Check if already processed in CRM
                            existing_lead = crm.get_lead_by_name(biz_name) if hasattr(crm, 'get_lead_by_name') else None
                            if existing_lead and existing_lead.get("status") in ["emailed", "form_submitted"]:
                                print(f"  ↪ Already contacted previously. Skipping.")
                                continue

                            lead_record = {
                                "name": biz_name,
                                "phone": phone,
                                "rating": rating_str,
                                "review_count": review_count_str,
                                "website": website,
                                "niche": niche,
                                "city": city,
                                "status": "discovered"
                            }
                            lead_id = crm.save_lead(lead_record)
                            lead_record["id"] = lead_id

                            if not website:
                                print(f"  ↪ No website listed. Logged to CRM as priority prospect.")
                                log_to_excel(biz_name, niche, city, phone, rating_str, review_count_str, "", "", "no_website", "High opportunity")
                                continue

                            # Open website in a lightweight tab with image/font blocking for speed & low RAM
                            site_page = context.new_page()
                            site_page.set_default_timeout(15000)

                            def block_heavy_assets(route):
                                if route.request.resource_type in ["image", "media", "font"]:
                                    route.abort()
                                else:
                                    route.continue_()

                            site_page.route("**/*", block_heavy_assets)

                            action_taken = False
                            try:
                                site_page.goto(website, timeout=15000, wait_until="domcontentloaded")
                                time.sleep(1.5)

                                # Attempt 1: Contact Form Submission
                                print(f"  ↪ Scanning for contact form on {website}...")
                                form_res = form_submitter.submit_contact_form(site_page, biz_name)

                                if form_res.get("status") == "submitted":
                                    new_count = increment_today_action_count()
                                    action_taken = True
                                    print(f"  ✅ Form successfully submitted! ({new_count}/{max_daily_limit} today)")
                                    crm.update_lead_status(lead_id, "form_submitted")
                                    log_to_excel(biz_name, niche, city, phone, rating_str, review_count_str, website, "", "form_submitted", "Verified submit")
                                    notifier.notify_form_submitted(biz_name, website, niche, city, new_count, max_daily_limit)

                                elif form_res.get("status") == "captcha":
                                    print(f"  ⚠️ CAPTCHA detected on form! Alerting Discord...")
                                    crm.update_lead_status(lead_id, "captcha_encountered")
                                    log_to_excel(biz_name, niche, city, phone, rating_str, review_count_str, website, "", "captcha_encountered", "Needs manual solve")
                                    notifier.notify_captcha_detected(biz_name, form_res.get("url", website), form_res.get("screenshot"))

                                else:
                                    print(f"  ↪ No suitable form or submit failed ({form_res.get('details')}). Falling back to email finder...")
                                    # Attempt 2: Direct Email Fallback
                                    found_email = email_finder.find_email_for_website(website)
                                    if found_email:
                                        print(f"  📧 Found email: {found_email}. Sending outreach email...")
                                        lead_record["email"] = found_email
                                        sent = email_sender.send_outreach_email(lead_record)
                                        if sent:
                                            new_count = increment_today_action_count()
                                            action_taken = True
                                            print(f"  ✅ Email sent to {found_email}! ({new_count}/{max_daily_limit} today)")
                                            crm.update_lead_status(lead_id, "emailed")
                                            log_to_excel(biz_name, niche, city, phone, rating_str, review_count_str, website, found_email, "emailed", "Direct email")
                                            notifier.notify_email_sent(biz_name, found_email, niche, city, new_count, max_daily_limit)
                                    else:
                                        print(f"  ↪ No direct email found on site.")
                                        log_to_excel(biz_name, niche, city, phone, rating_str, review_count_str, website, "", "no_form_or_email", "")

                            except Exception as e:
                                print(f"  ❌ Error visiting {website}: {e}")
                                log_to_excel(biz_name, niche, city, phone, rating_str, review_count_str, website, "", "site_error", str(e))
                            finally:
                                site_page.close()

                            if action_taken:
                                sleep_time = delay_seconds + random.uniform(2.0, 4.0)
                                print(f"  ⏳ Pausing {sleep_time:.1f}s before next business...")
                                time.sleep(sleep_time)

                        except Exception as e:
                            print(f"  ❌ Error processing card {idx+1}: {e}")
                            continue

                except Exception as e:
                    print(f"❌ Error during search '{query}': {e}")
                    notifier.notify_error(query, str(e))
                    continue

        browser.close()
        print("\n🎉 Completed all targets for this run!")

if __name__ == "__main__":
    run_maps_autonomous(headless=True)
